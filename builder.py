"""Сборка файла презентации из плана слайдов (нейтральный шаблон).

Все внешние параметры оформления собраны в блоке TEMPLATE.
Чтобы подогнать под вуз: поменяйте цвета, шрифт, подпись и (по желанию) логотип.
"""
import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

# ---------- НАСТРОЙКИ ШАБЛОНА (меняйте здесь, только то, что внутри кавычек) ----------
TEMPLATE = {
    "font": "Calibri",             # шрифт (должен быть установлен на компьютере)
    "main_color": "1F3A5F",        # основной цвет: заголовки, титульный слайд
    "accent_color": "C9A227",      # акцентная полоса
    "text_color": "222222",        # цвет основного текста
    "light_text": "FFFFFF",        # текст на тёмном фоне
    "footer_text": "",             # подпись внизу слайдов, например название вуза
    "logo_path": "",               # имя файла логотипа в папке проекта, например "logo.png"
}
# ------------------------------------------------------------------------------------

W, H = Inches(13.333), Inches(7.5)
BODY_W_IN, BODY_H_IN = 11.5, 4.5      # область под текст пунктов, дюймы
SPACE_AFTER_PT = 12                   # отступ между пунктами
FONT_SIZES = [24, 22, 20, 18, 16]     # какие размеры можно подбирать для пунктов
MAX_BULLET_CHARS = 400
MAX_TITLE_CHARS = 160


# ---------- Очистка текста ----------
def clean_text(text):
    """Одна строка: убирает переносы, лишние пробелы и символы, ломающие файл."""
    text = str(text).encode("utf-8", "ignore").decode("utf-8", "ignore")
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[\x00-\x08\x0e-\x1f\x7f\ufffe\uffff]", "", text)
    return text.strip()


def clean_content(text):
    """Многострочный текст: сохраняет переносы строк, убирает опасные символы."""
    text = str(text).encode("utf-8", "ignore").decode("utf-8", "ignore")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ufffe\uffff]", "", text)
    return text.strip()


def shorten(text, limit):
    """Обрезает по границе слова и добавляет многоточие."""
    if len(text) <= limit:
        return text
    cut = text[:limit - 1]
    if " " in cut[limit // 2:]:
        cut = cut[:cut.rfind(" ")]
    return cut.rstrip(" ,;:-") + "…"


# ---------- Рисование ----------
def rgb(hex_color):
    value = str(hex_color).strip().lstrip("#")
    try:
        return RGBColor.from_string(value)
    except ValueError:
        raise ValueError(f"Неверный цвет «{hex_color}» в блоке TEMPLATE. Нужно 6 символов, например 1F3A5F")


def add_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_text(slide, x, y, w, h, text, size, color, bold=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = TEMPLATE["font"]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return frame


def add_logo(slide, x, y, height):
    path = TEMPLATE["logo_path"]
    if not path:
        return
    logo = Path(path)
    if not logo.is_absolute():
        logo = Path(__file__).resolve().parent / logo
    if logo.exists():
        try:
            slide.shapes.add_picture(str(logo), x, y, height=height)
        except Exception:
            pass  # битый файл логотипа - просто пропускаем


# ---------- Подбор размера текста ----------
def estimate_height_pt(bullets, size):
    chars_per_line = max(10, int(BODY_W_IN * 72 / (size * 0.55)))
    lines = sum(max(1, math.ceil((len(b) + 3) / chars_per_line)) for b in bullets)
    return lines * size * 1.2 + len(bullets) * SPACE_AFTER_PT


def choose_size(bullets):
    """Самый крупный размер, при котором пункты помещаются; None, если не помещаются."""
    for size in FONT_SIZES:
        if estimate_height_pt(bullets, size) <= BODY_H_IN * 72 * 0.95:
            return size
    return None


def split_bullets(bullets):
    """Делит пункты на части, чтобы каждая часть помещалась на одном слайде."""
    chunks, current = [], []
    for bullet in bullets:
        if current and choose_size(current + [bullet]) is None:
            chunks.append(current)
            current = [bullet]
        else:
            current.append(bullet)
    if current:
        chunks.append(current)
    return chunks


# ---------- Слайды ----------
def title_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # пустой макет
    add_rect(slide, 0, 0, W, H, TEMPLATE["main_color"])
    add_rect(slide, Inches(0.9), Inches(3.0), Inches(1.6), Inches(0.12), TEMPLATE["accent_color"])
    n = len(title)
    size = 44 if n <= 40 else 36 if n <= 80 else 30 if n <= 120 else 26
    add_text(slide, Inches(0.9), Inches(3.3), Inches(11.5), Inches(2.6),
             title, size, TEMPLATE["light_text"], bold=True)
    if TEMPLATE["footer_text"]:
        add_text(slide, Inches(0.9), Inches(6.5), Inches(11), Inches(0.5),
                 clean_text(TEMPLATE["footer_text"]), 16, TEMPLATE["light_text"])
    add_logo(slide, Inches(0.9), Inches(0.7), Inches(1.0))


def content_slide(prs, title, bullets, number, size=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, Inches(0.25), H, TEMPLATE["main_color"])
    n = len(title)
    title_size = 32 if n <= 45 else 28 if n <= 80 else 24 if n <= 120 else 20
    add_text(slide, Inches(0.9), Inches(0.5), Inches(10.8), Inches(1.2),
             title, title_size, TEMPLATE["main_color"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(0.9), Inches(1.75), Inches(1.2), Inches(0.07), TEMPLATE["accent_color"])

    if bullets:
        size = size or choose_size(bullets) or FONT_SIZES[-1]
        box = slide.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(BODY_W_IN), Inches(BODY_H_IN))
        frame = box.text_frame
        frame.word_wrap = True
        for i, bullet in enumerate(bullets):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = p.add_run()
            run.text = "•  " + bullet
            run.font.name = TEMPLATE["font"]
            run.font.size = Pt(size)
            run.font.color.rgb = rgb(TEMPLATE["text_color"])
            p.space_after = Pt(SPACE_AFTER_PT)

    if TEMPLATE["footer_text"]:
        add_text(slide, Inches(0.9), Inches(6.9), Inches(9), Inches(0.4),
                 clean_text(TEMPLATE["footer_text"]), 12, "777777")
    add_text(slide, Inches(11.6), Inches(6.9), Inches(1.2), Inches(0.4), str(number), 12, "777777")
    add_logo(slide, Inches(11.9), Inches(0.5), Inches(0.7))


def build_presentation(topic, outline, out_path):
    """outline - список слайдов: [{"title": "...", "bullets": ["...", ...]}, ...]
    Первый слайд всегда титульный. Если пунктов слишком много, слайд делится на части."""
    if not outline:
        raise ValueError("Пустой план презентации")

    topic = shorten(clean_text(topic), MAX_TITLE_CHARS) or "Презентация"
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    prs.core_properties.title = topic
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""

    number = 0
    for i, slide_data in enumerate(outline):
        title = shorten(clean_text(slide_data.get("title") or topic), MAX_TITLE_CHARS)
        bullets = [shorten(clean_text(b), MAX_BULLET_CHARS) for b in (slide_data.get("bullets") or [])]
        bullets = [b for b in bullets if b]
        if i == 0:
            title_slide(prs, title)
            number += 1
            continue
        parts = split_bullets(bullets) or [[]]
        # у всех частей одного слайда размер шрифта одинаковый
        sizes = [choose_size(part) or FONT_SIZES[-1] for part in parts if part]
        common_size = min(sizes) if sizes else None
        for part_index, part in enumerate(parts):
            part_title = title if part_index == 0 else shorten(title, 100) + " (продолжение)"
            number += 1
            content_slide(prs, part_title, part, number, common_size)

    prs.save(str(out_path))
    return out_path
